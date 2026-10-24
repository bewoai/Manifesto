"""Rezervasyon birleşimleri arkadaşın düzenini izler: PAX/otel/pickup/rezerve/acente birleşik,
irtibat (E) ve not (M) yolcu başına. Taşıma bloğun kendi birleşimlerini korur."""
from __future__ import annotations

import openpyxl

from app import config
from app.manifest.excel_cache import clear_excel_cache
from app.manifest.planning import create_reservation, reorder_reservation

CORE = (config.COL_PAX, config.COL_HOTEL, config.COL_PICKUP, config.COL_RESERVED_BY, config.COL_AGENCY)


def _ranges(path, sheet="24.10.2026"):
    ws = openpyxl.load_workbook(path)[sheet]
    return {(m.min_col, m.min_row, m.max_row) for m in ws.merged_cells.ranges}


def _two_blocks(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "24.10.2026"
    ws.cell(config.PLANNING_HEADER_ROW, config.COL_PAX).value = "PAX"
    for lead, pax, hotel in ((4, 3, "ARGOS"), (7, 2, "KELEBEK")):
        ws.cell(lead, config.COL_PAX).value = pax
        ws.cell(lead, config.COL_HOTEL).value = hotel
        for r in range(lead, lead + pax):
            ws.cell(r, config.COL_NAME).value = f"KISI {r}"
            ws.cell(r, config.COL_BALLOON).value = "BYF"
        for c in CORE:
            ws.merge_cells(start_row=lead, start_column=c, end_row=lead + pax - 1, end_column=c)
    ws.cell(4, config.COL_ROOM).value = "0 532"            # ilk blokta irtibat birleşik (arkadaş öyle yazmış)
    ws.merge_cells(start_row=4, start_column=config.COL_ROOM, end_row=6, end_column=config.COL_ROOM)
    wb.save(path)


def test_new_reservation_does_not_merge_room_or_note(tmp_path):
    p = tmp_path / "plan.xlsx"
    _two_blocks(p)
    create_reservation(p, "24.10.2026", pax=3, balloon="BYF", fields={"hotel": "ASMALI", "note": "VEGAN"})
    ranges = _ranges(p)
    lead = 9
    for c in CORE:
        assert (c, lead, lead + 2) in ranges
    assert not any(c in (config.COL_ROOM, config.COL_NOTE) and lo == lead for c, lo, _ in ranges)
    assert openpyxl.load_workbook(p)["24.10.2026"].cell(lead, config.COL_NOTE).value == "VEGAN"


def test_reorder_keeps_block_own_merges(tmp_path):
    p = tmp_path / "plan.xlsx"
    _two_blocks(p)
    clear_excel_cache()
    reorder_reservation(p, "24.10.2026", [7, 8], 4)       # KELEBEK bloğunu en üste taşı
    ranges = _ranges(p)
    for c in CORE:
        assert (c, 4, 5) in ranges and (c, 6, 8) in ranges
    assert (config.COL_ROOM, 6, 8) in ranges               # ARGOS'un irtibat birleşimi onunla kaydı
    assert not any(c == config.COL_NOTE for c, _, _ in ranges)
    assert not any(c == config.COL_ROOM and lo == 4 for c, lo, _ in ranges)
    ws = openpyxl.load_workbook(p)["24.10.2026"]
    assert ws.cell(4, config.COL_HOTEL).value == "KELEBEK" and ws.cell(6, config.COL_HOTEL).value == "ARGOS"
