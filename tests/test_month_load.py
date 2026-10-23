"""Takvim doluluğu: /api/planning/month-load (koltuk şablonu fixture'ı, PII yok).

  21.10 tam dolu · 22.10 kısmen dolu (BYF bandı boş) · 23.10 az dolu — her gün 4 bant × 28.
"""
from __future__ import annotations

import importlib
import shutil
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import config
from app.manifest.excel_cache import clear_excel_cache
from app.manifest.planning import month_load

FIXTURE = Path(__file__).parent / "fixtures" / "seat_template_ekim.xlsx"
CODES = ["BYF", "BTK", "BYJ", "BZR", "BZV"]
EXPECTED = {
    "21.10.2026": {"used": 112, "free": 0, "total": 112},
    "22.10.2026": {"used": 70, "free": 42, "total": 112},
    "23.10.2026": {"used": 40, "free": 72, "total": 112},
}


@pytest.fixture()
def wb_path(tmp_path):
    p = tmp_path / "ekim.xlsx"
    shutil.copy2(FIXTURE, p)
    clear_excel_cache()
    return p


@pytest.fixture()
def client(tmp_path, monkeypatch, wb_path):
    monkeypatch.setenv("MANIFESTO_SETTINGS", str(tmp_path / "settings.json"))
    monkeypatch.setenv("IRTIFA_DB_PATH", str(tmp_path / "month.db"))
    import app.settings as settings_mod
    importlib.reload(settings_mod)
    import app.main as main_mod
    importlib.reload(main_mod)
    c = TestClient(main_mod.app)
    c.post("/api/auth/setup", json={"username": "admin", "display_name": "Y", "password": "strong-pass"})
    s = settings_mod.load()
    s.planning_xlsx = str(wb_path); s.balloon_codes = list(CODES); s.is_setup_complete = True
    settings_mod.save(s)
    return c


def test_month_load_seat_template(wb_path):
    assert month_load(wb_path, 2026, 10, capacity=28, codes=CODES) == EXPECTED
    assert month_load(wb_path, 2026, 11, capacity=28, codes=CODES) == {}


def test_month_load_append_template_uses_roster_capacity(wb_path):
    # Ekleme şablonu (bant yok): toplam = kapasite × (Ayarlar kodları ∪ sayfadaki balonlar)
    wb = openpyxl.load_workbook(wb_path)
    ws = wb.create_sheet("24.10.2026")
    first = config.PLANNING_FIRST_DATA_ROW
    ws.cell(first, config.COL_PAX, 3)
    for i, name in enumerate(["YOLCU A", "YOLCU B", "YOLCU C"]):
        ws.cell(first + i, config.COL_NAME, name)
        ws.cell(first + i, config.COL_BALLOON, "BYF")
    wb.save(wb_path)
    wb.close()
    day = month_load(wb_path, 2026, 10, capacity=28, codes=CODES)["24.10.2026"]
    assert day == {"used": 3, "free": 0, "total": 5 * 28}


def test_api_month_load_is_cached_by_revision(client, wb_path):
    first = client.get("/api/planning/month-load", params={"month": "2026-10"})
    assert first.status_code == 200, first.text
    body = first.json()
    assert body["days"] == EXPECTED
    assert body["cached"] is False

    second = client.get("/api/planning/month-load", params={"month": "2026-10"}).json()
    assert second["cached"] is True
    assert second["workbook_revision"] == body["workbook_revision"]
    assert second["days"] == EXPECTED

    # Excel değişti → yeni revizyon, yeniden okunur (tek kaynak Excel)
    load = client.get("/api/planning/load", params={"sheet": "22.10.2026"}).json()
    r = client.post("/api/planning/create-with-identities", json={
        "sheet": "22.10.2026", "pax": 2, "hotel": "OTEL 5", "agency": "ACENTE 5",
        "identities": [{"nationality": "KOR", "sex": "M", "name": "YOLCU A", "passport_no": "M12345678"},
                       {"nationality": "KOR", "sex": "F", "name": "YOLCU B", "passport_no": "M87654321"}],
        "expected_revision": load["workbook_revision"],
    })
    assert r.status_code == 200, r.text
    third = client.get("/api/planning/month-load", params={"month": "2026-10"}).json()
    assert third["cached"] is False
    assert third["workbook_revision"] != body["workbook_revision"]
    assert third["days"]["22.10.2026"] == {"used": 72, "free": 40, "total": 112}


@pytest.mark.parametrize("month", ["2026-13", "2026-1", "ekim"])
def test_api_month_load_rejects_bad_month(client, month):
    assert client.get("/api/planning/month-load", params={"month": month}).status_code == 422
