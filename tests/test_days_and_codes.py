"""Boş gün oluşturma (koltuk düzeni korunur), uyruk kodu eşanlamları, elle kimlik düzeltme."""
from __future__ import annotations

import datetime as dt
import shutil
from pathlib import Path

import openpyxl
import pytest

from app import config
from app.country.country_map import iso3_to_name, normalize_code
from app.manifest.excel_cache import clear_excel_cache
from app.manifest.planning import (
    create_day_sheet, create_missing_days, free_seat_summary, is_seat_layout, nearest_source_sheet,
    read_blocks, read_rows, write_identity,
)
from app.manifest.writer import build_rows

FIXTURE = Path(__file__).parent / "fixtures" / "seat_template_ekim.xlsx"


@pytest.fixture()
def wb_path(tmp_path):
    p = tmp_path / "ekim.xlsx"
    shutil.copy2(FIXTURE, p)
    clear_excel_cache()
    return p


# ─── Boş gün ─────────────────────────────────────────────────────────────

def test_nearest_source_prefers_previous_day():
    names = ["01.10.2026", "04.10.2026", "24.10.2026", "Sayfa1"]
    assert nearest_source_sheet(names, "25.10.2026") == "24.10.2026"
    assert nearest_source_sheet(names, "02.10.2026") == "01.10.2026"
    assert nearest_source_sheet(names, "30.09.2026") == "01.10.2026"     # öncesi yoksa en yakın sonraki
    assert nearest_source_sheet(["Sayfa1"], "25.10.2026") is None


def test_blank_day_keeps_seat_layout_and_clears_reservations(wb_path):
    create_day_sheet(wb_path, "24.10.2026")
    clear_excel_cache()
    wb = openpyxl.load_workbook(wb_path)
    assert wb.sheetnames == ["21.10.2026", "22.10.2026", "23.10.2026", "24.10.2026"]
    ws = wb["24.10.2026"]
    assert is_seat_layout(wb["23.10.2026"])
    assert not [r for r in ws.merged_cells.ranges if r.min_row >= config.PLANNING_FIRST_DATA_ROW]
    assert any(ws.cell(1, c).value == dt.datetime(2026, 10, 24) for c in range(1, 15))
    rows, blocks = read_rows(wb_path, "24.10.2026"), read_blocks(wb_path, "24.10.2026")
    assert blocks == [] and not any(r.name or r.passport_no for r in rows)
    seats = free_seat_summary(rows, blocks)
    assert sum(v["free"] for v in seats.values()) == 112                 # 23.10'un bant düzeni
    assert all(ws.cell(r.row, config.COL_COMPANY).value == "THK" for r in rows)


def test_create_missing_days_in_order_and_idempotent(wb_path):
    created = create_missing_days(wb_path, ["25.10.2026", "22.10.2026", "24.10.2026", "x"])
    assert created == ["24.10.2026", "25.10.2026"]
    assert create_missing_days(wb_path, ["24.10.2026"]) == []
    names = openpyxl.load_workbook(wb_path, read_only=True).sheetnames
    assert names == ["21.10.2026", "22.10.2026", "23.10.2026", "24.10.2026", "25.10.2026"]


def test_blank_day_from_append_template_clears_everything(tmp_path):
    p = tmp_path / "plan.xlsx"
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "01.11.2026"
    ws.cell(4, config.COL_PAX).value = 2; ws.cell(4, config.COL_NAME).value = "A"; ws.cell(4, config.COL_BALLOON).value = "BYF"
    ws.cell(5, config.COL_NAME).value = "B"; ws.cell(5, config.COL_BALLOON).value = "BYF"
    wb.save(p); clear_excel_cache()
    create_day_sheet(p, "02.11.2026")
    ws2 = openpyxl.load_workbook(p)["02.11.2026"]
    assert ws2.cell(4, config.COL_BALLOON).value is None and ws2.cell(4, config.COL_NAME).value is None


# ─── Uyruk kodları ──────────────────────────────────────────────────────

@pytest.mark.parametrize("raw,iso", [
    ("IRE", "IRL"), ("ire", "IRL"), ("IRL", "IRL"), ("UK", "GBR"), ("GB", "GBR"), ("ENG", "GBR"),
    ("D", "DEU"), ("D<<", "DEU"), ("GER", "DEU"), ("SUI", "CHE"), ("NED", "NLD"), ("Türkiye", "TUR"),
    ("T.C.", "TUR"), ("Ireland", "IRL"), ("united kingdom", "GBR"), ("İngiltere", "GBR"),
])
def test_normalize_code(raw, iso):
    assert normalize_code(raw) == iso


def test_unknown_and_ambiguous_codes_untouched():
    assert normalize_code("ZZZ") == "ZZZ"
    assert normalize_code("CHI") == "CHI"       # Şili mi Çin mi belirsiz → dokunulmaz
    assert normalize_code("") == ""
    assert iso3_to_name("IRE") == "Ireland"
    assert iso3_to_name("ZZZ") is None


def test_manifest_resolves_ire_without_warning(wb_path):
    rows = read_rows(wb_path, "22.10.2026")
    target = next(r for r in rows if r.balloon == "BZR" and r.name)
    write_identity(wb_path, "22.10.2026", {target.row: {"nationality": "IRE"}})
    clear_excel_cache()
    man = {m.name: m for m in build_rows(wb_path, "22.10.2026", "BZR")}
    m = man[target.name]
    assert m.nationality == "Ireland" and not m.warnings


# ─── Elle kimlik düzeltme (API normalizasyonu) ──────────────────────────

def test_manual_identity_edit_is_normalized():
    from app.main import _normalize_identity_fields
    out = _normalize_identity_fields({"nationality": "ire", "sex": "Kadın", "name": "  esther   adebayo ",
                                      "passport_no": "pr 869 5163"})
    assert out == {"nationality": "IRL", "sex": "F", "name": "ESTHER ADEBAYO", "passport_no": "PR8695163"}
    assert _normalize_identity_fields({"sex": "E"}) == {"sex": "M"}
    assert _normalize_identity_fields({}) == {}
