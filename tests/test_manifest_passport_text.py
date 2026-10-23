"""Manifestoda pasaport sütunu metin olmalı (SHGM sistemi sayı biçimini bazen reddediyor)."""
from __future__ import annotations

import openpyxl
import pytest

from app import config
from app.manifest.writer import ManifestRow, passport_text, write_manifest


@pytest.mark.parametrize("raw,text", [
    ("U1234567", "U1234567"), ("01234567", "01234567"), (1234567, "1234567"),
    (1234567.0, "1234567"), ("1234567.0", "1234567"), (" 98765 ", "98765"), (None, ""),
])
def test_passport_text(raw, text):
    assert passport_text(raw) == text


def test_passport_column_is_text_formatted(tmp_path):
    rows = [
        ManifestRow(name="A B", sex="Erkek", nationality="Turkey", passport_no="01234567", warnings=[]),
        ManifestRow(name="C D", sex="Kadın", nationality="Italy", passport_no="U7654321", warnings=[]),
    ]
    out = write_manifest(rows, tmp_path / "BZR.xlsx")
    ws = openpyxl.load_workbook(out)[config.MANIFEST_SHEET]
    first = config.MANIFEST_FIRST_DATA_ROW
    for i, expected in enumerate(["01234567", "U7654321"]):
        cell = ws.cell(first + i, config.M_COL_PASSPORT)
        assert cell.value == expected          # baştaki sıfır korunur
        assert cell.data_type == "s"           # sayı değil metin hücresi
        assert cell.number_format == "@"       # Excel'de "Metin" biçimi
    # Şablonda boş kalan satırlar da metin biçiminde (elle eklenen numara sayıya dönmesin)
    assert ws.cell(first + 5, config.M_COL_PASSPORT).number_format == "@"
