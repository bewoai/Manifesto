"""Otel kartları: yazım birleştirme, bölge, buluşma noktası; gruplamaya etkisi."""
from __future__ import annotations

import importlib
import json

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import config
from app.hotels import (
    HotelCard, HotelResolver, delete_card, load_catalog, match_key, merge_names, similar, unmatched_groups, upsert_card,
)
from app.manifest.excel_cache import clear_excel_cache
from app.manifest.planning import ReservationBlock, hotel_seats_by_balloon, hotel_split


def test_match_key_and_similarity():
    assert match_key("Argos in Cappadocia") == "ARGOS"
    assert match_key("Göreme Inn") == "INN"            # yer adı atılır
    assert similar("MARRIOT", "MARRIOTT")
    assert similar("CROWN PLAZA", "CROWNE PLAZA")
    assert similar("ROCCA", "ROCCA STONE HOUSE")
    assert not similar("GOREME INN", "GOREME PALACE CAVE SUITES")
    assert not similar("KEMAL STONE HOUSE", "ROCCA STONE HOUSE")


def test_unmatched_groups_lead_by_most_frequent_and_skip_carded():
    counts = {"MARRIOTT": 5, "Marriot": 2, "marriott": 1, "ARGOS": 3, "Argos in Cappadocia": 1,
              "KELEBEK": 2, "STONE HOUSE CAVE": 1, "ROCCA STONE HOUSE": 1}
    groups = unmatched_groups(counts, [])
    first = groups[0]
    assert first["suggested_name"] == "MARRIOTT" and first["total"] == 8
    assert [n["name"] for n in first["names"]] == ["MARRIOTT", "Marriot"]   # aynı katlanmış ad tek yazım
    names = [[n["name"] for n in g["names"]] for g in groups]
    assert ["ARGOS", "Argos in Cappadocia"] in names
    assert ["STONE HOUSE CAVE"] in names and ["ROCCA STONE HOUSE"] in names

    cards = [HotelCard(id="a1", name="ARGOS CAVE", aliases=["ARGOS"])]
    groups = unmatched_groups(counts, cards)
    argos = next(g for g in groups if any(n["name"] == "Argos in Cappadocia" for n in g["names"]))
    assert argos["card_id"] == "a1"                          # benzer kart önerilir
    assert not any(n["name"] == "ARGOS" for g in groups for n in g["names"])   # karta bağlı, listede yok


def test_merge_upsert_delete_keep_one_card_per_name(tmp_path):
    path = tmp_path / "hotels.json"
    a = merge_names(["MARRIOTT", "Marriot"], name="Marriott Cappadocia", region="ürgüp", path=path)
    assert a.region == "Ürgüp" and {x.upper() for x in a.aliases} == {"MARRIOTT", "MARRIOT"}
    b = merge_names(["KELEBEK", "Kelebek Cave"], path=path)
    assert b.name == "KELEBEK" and b.aliases == ["Kelebek Cave"]
    # yazım başka karta taşınınca eskisinden düşer
    merge_names(["Marriot"], card_id=b.id, path=path)
    cards = {c.id: c for c in load_catalog(path)}
    assert "Marriot" not in cards[a.id].aliases and "Marriot" in cards[b.id].aliases
    # ad değişince eski ad yazım olarak kalır
    c = upsert_card({"id": b.id, "name": "KELEBEK SPECIAL CAVE", "meeting_point": "Kapı önü",
                     "map_url": "https://maps.example/x"}, path=path)
    res = HotelResolver(load_catalog(path))
    assert res.resolve("kelebek").id == c.id and res.display("KELEBEK") == "KELEBEK SPECIAL CAVE"
    assert res.resolve("kelebek").meeting_point == "Kapı önü"
    assert res.key("Kelebek Cave") == res.key("KELEBEK") == f"#{c.id}"
    assert res.key("YOK OTEL") == "YOK OTEL"
    assert upsert_card({"id": c.id, "name": c.name, "map_url": "javascript:alert(1)"}, path=path).map_url == ""
    with pytest.raises(ValueError):
        upsert_card({"name": "  "}, path=path)
    with pytest.raises(KeyError):
        merge_names(["X"], card_id="yok", path=path)
    assert delete_card(a.id, path=path) and not delete_card(a.id, path=path)
    assert json.loads(path.read_text(encoding="utf-8"))["hotels"][0]["name"] == "KELEBEK SPECIAL CAVE"


def _blk(row, hotel, balloon, pax=2, pickup="05:00"):
    return ReservationBlock(lead_row=row, rows=list(range(row, row + pax)), pax=pax, agency="A",
                            balloon=balloon, hotel=hotel, pickup=pickup, room="", reserved_by="",
                            company="THK", pilot="", note="", driver="", coming_place="", lead_name="")


def test_grouping_uses_card_key():
    blocks = [_blk(4, "MARRIOTT", "BYF"), _blk(6, "Marriot", "BYJ", pax=3)]
    assert hotel_split(blocks) == []                      # kartsız: iki ayrı otel
    res = HotelResolver([HotelCard(id="m1", name="MARRIOTT", aliases=["Marriot"])])
    split = hotel_split(blocks, res.key)
    assert len(split) == 1 and split[0]["balloons"] == ["BYJ", "BYF"] and split[0]["seats"] == 5
    assert hotel_seats_by_balloon(blocks, "marriot", "05:00", res.key) == {"BYJ": 3, "BYF": 2}


@pytest.fixture()
def client(monkeypatch, tmp_path):
    plan = tmp_path / "plan.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "24.10.2026"
    for lead, pax, hotel in ((4, 2, "MARRIOTT"), (6, 1, "Marriot"), (7, 1, "KELEBEK")):
        ws.cell(lead, config.COL_PAX).value = pax
        ws.cell(lead, config.COL_HOTEL).value = hotel
        for r in range(lead, lead + pax):
            ws.cell(r, config.COL_NAME).value = f"K{r}"
            ws.cell(r, config.COL_BALLOON).value = "BYF"
    wb.save(plan)
    settings_path = tmp_path / "settings.json"
    settings_path.write_text(json.dumps({"planning_xlsx": str(plan)}), encoding="utf-8")
    monkeypatch.setenv("MANIFESTO_SETTINGS", str(settings_path))
    monkeypatch.setenv("IRTIFA_DB_PATH", str(tmp_path / "m.db"))
    monkeypatch.delenv("IRTIFA_HOTELS_PATH", raising=False)
    import app.settings as settings_mod
    importlib.reload(settings_mod)
    import app.main as main_mod
    importlib.reload(main_mod)
    clear_excel_cache()
    with TestClient(main_mod.app) as c:
        yield c, tmp_path


def test_api_flow(client):
    c, tmp = client
    data = c.get("/api/hotels").json()
    assert data["stats"] == {"reservations": 3, "matched": 0, "cards": 0, "with_region": 0}
    assert "Göreme" in data["regions"]
    top = data["unmatched"][0]
    assert {n["name"] for n in top["names"]} == {"MARRIOTT", "Marriot"} and top["total"] == 2

    r = c.post("/api/hotels/merge", json={"names": ["MARRIOTT", "Marriot"], "name": "MARRIOTT", "region": "Ürgüp"})
    assert r.status_code == 200, r.text
    card = r.json()["card"]
    assert (tmp / "hotels.json").exists()                  # ayarlarla aynı klasörde

    r = c.post("/api/hotels/card", json={"id": card["id"], "name": "MARRIOTT", "meeting_point": "Lobi"})
    assert r.status_code == 200 and r.json()["card"]["meeting_point"] == "Lobi"

    data = c.get("/api/hotels").json()
    assert data["stats"]["matched"] == 2 and data["stats"]["with_region"] == 1
    assert data["cards"][0]["count"] == 2
    assert [g["suggested_name"] for g in data["unmatched"]] == ["KELEBEK"]

    assert c.post("/api/hotels/merge", json={"names": []}).status_code == 422
    assert c.post("/api/hotels/card", json={"name": ""}).status_code == 422
    assert c.post("/api/hotels/delete", json={"id": "yok"}).status_code == 404
    assert c.post("/api/hotels/delete", json={"id": card["id"]}).status_code == 200
    assert c.get("/api/hotels").json()["stats"]["cards"] == 0
