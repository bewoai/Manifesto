"""Şoför listeleri: rezervasyon bazlı gruplama (birleşik hücre / rowspan), NOT, HH:MM."""
import datetime
import sys
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from app import config
from app.manifest.planning import read_blocks
from app.reports import (
    HEADERS,
    LAPIS,
    _pdf_table_parts,
    _register_pdf_fonts,
    driver_groups,
    driver_rows,
    export_driver_reports,
    format_pickup,
)

SHEET = "22.10.2026"
NOTE = "Bebek koltuğu, çıkışta arayın"


def _write_block(ws, start, *, pax, names, driver, pickup, note="", hotel="ARGOS",
                 balloon="BZR", agency="ACME", room="ODA 202", coming="DİREKT ALAN"):
    for offset, name in enumerate(names):
        r = start + offset
        ws.cell(r, config.COL_NAME).value = name
        ws.cell(r, config.COL_COMPANY).value = "THK"
        ws.cell(r, config.COL_BALLOON).value = balloon
        ws.cell(r, config.COL_DRIVER).value = driver
        ws.cell(r, config.COL_COMING_PLACE).value = coming
    ws.cell(start, config.COL_PAX).value = pax
    ws.cell(start, config.COL_HOTEL).value = hotel
    ws.cell(start, config.COL_ROOM).value = room
    ws.cell(start, config.COL_PICKUP).value = pickup
    ws.cell(start, config.COL_AGENCY).value = agency
    if note:
        ws.cell(start, config.COL_NOTE).value = note
    return start + len(names)


def _planning(path: Path, *, big: int = 0) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.cell(config.PLANNING_HEADER_ROW, config.COL_PAX).value = "PAX"
    ws.cell(config.PLANNING_HEADER_ROW, config.COL_NAME).value = "İSİM"
    ws.cell(config.PLANNING_HEADER_ROW, config.COL_NOTE).value = "NOT"
    row = config.PLANNING_HEADER_ROW + 1
    # Şoför 2: 05:20'de 3 kişilik (notlu), 04:50'de tek kişilik rezervasyon.
    row = _write_block(ws, row, pax=3, names=["ALİ VELİ", "AYŞE VELİ", "ÖZGE VELİ"],
                       driver="2", pickup=datetime.time(5, 20), note=NOTE)
    row = _write_block(ws, row, pax=1, names=["ŞÜKRÜ GÜL"], driver="2",
                       pickup="04:50:00", hotel="KAYA")
    # Şoför 1: datetime olarak saklanmış pickup.
    row = _write_block(ws, row, pax=2, names=["JOHN DOE", "JANE DOE"], driver="1",
                       pickup=datetime.datetime(1900, 1, 1, 6, 5), balloon="BYF")
    if big:
        names = [f"YOLCU {i:02d}" for i in range(big)]
        _write_block(ws, row, pax=big, names=names, driver="3", pickup="05:00",
                     note="Büyük grup")
    wb.save(path)
    wb.close()
    return path


@pytest.mark.parametrize("value, expected", [
    (datetime.time(5, 20), "05:20"),
    (datetime.time(5, 20, 59), "05:20"),
    (datetime.datetime(1900, 1, 1, 5, 20), "05:20"),
    ("05:20:00", "05:20"),
    ("5:20", "05:20"),
    ("1900-01-01 05:20:00", "05:20"),
    ("2026-06-22T04:10:00", "04:10"),
    ("04.10", "04:10"),
    ("0.2222222222", "05:20"),
    ("", ""),
    (None, ""),
    ("SABAH", "SABAH"),
])
def test_format_pickup(value, expected):
    assert format_pickup(value) == expected


def test_groups_carry_note_and_hhmm(tmp_path):
    path = _planning(tmp_path / "planning.xlsx")
    groups = driver_groups(read_blocks(path, SHEET))
    assert list(groups) == ["1", "2"]
    first, second = groups["2"]
    # Pickup sırası: 04:50 önce.
    assert first["pickup"] == "04:50" and first["passengers"] == ["ŞÜKRÜ GÜL"]
    assert second["pickup"] == "05:20"
    assert second["note"] == NOTE
    assert second["passengers"] == ["ALİ VELİ", "AYŞE VELİ", "ÖZGE VELİ"]
    assert groups["1"][0]["pickup"] == "06:05"

    flat = driver_rows(read_blocks(path, SHEET))["2"]
    assert [r["pax"] for r in flat] == [1, 3, "", ""]
    assert all(r["note"] == NOTE for r in flat[1:])


def test_xlsx_merges_reservation_fields(tmp_path):
    path = _planning(tmp_path / "planning.xlsx")
    outputs = export_driver_reports(path, SHEET, tmp_path / "out", selected_driver="2")
    xlsx = next(p for p in outputs if p.suffix == ".xlsx")
    wb = openpyxl.load_workbook(xlsx)
    ws = wb.active
    assert ws["A1"].value == f"İRTİFA - {SHEET} - ŞOFÖR 2"

    header = [ws.cell(3, c).value for c in range(1, len(HEADERS) + 1)]
    assert header == HEADERS and header[-1] == "NOT"
    assert ws.cell(3, 1).fill.fgColor.rgb.endswith(LAPIS)

    # Satır 4: tek kişilik rezervasyon; 5-7: üç kişilik rezervasyon.
    assert ws["B4"].value == "ŞÜKRÜ GÜL" and ws["G4"].value == "04:50"
    assert [ws.cell(r, 2).value for r in (5, 6, 7)] == ["ALİ VELİ", "AYŞE VELİ", "ÖZGE VELİ"]
    assert [ws.cell(r, 1).value for r in (4, 5, 6, 7)] == [1, 2, 3, 4]
    assert ws["G5"].value == "05:20"
    assert ws["C5"].value == 3
    assert ws["K5"].value == NOTE

    merged = {str(rng) for rng in ws.merged_cells.ranges}
    for col in range(3, len(HEADERS) + 1):       # PAX .. NOT
        letter = get_column_letter(col)
        assert f"{letter}5:{letter}7" in merged
        assert ws.cell(5, col).alignment.vertical == "center"
    # Yolcu başına kolonlar ve tek kişilik rezervasyon birleştirilmez.
    assert not any(r.startswith(("A5", "B5")) or r.startswith(("C4", "K4")) for r in merged)
    # Rezervasyonlar arası kalın çizgi.
    assert ws["B5"].border.top.style == "medium"
    assert ws["B6"].border.top.style == "thin"
    wb.close()


def test_pdf_rowspans_and_note(tmp_path):
    path = _planning(tmp_path / "planning.xlsx")
    groups = driver_groups(read_blocks(path, SHEET))["2"]
    data, style = _pdf_table_parts(groups, "Helvetica", "Helvetica-Bold")
    assert [p.text for p in data[0]] == HEADERS
    assert len(data) == 1 + 4
    spans = {(cmd[1], cmd[2]) for cmd in style if cmd[0] == "SPAN"}
    # Tablo satırı 2-4: üç kişilik rezervasyon; kolon 2..10 (PAX..NOT) birleşik.
    assert spans == {((c, 2), (c, 4)) for c in range(2, len(HEADERS))}
    assert any(cmd[0] == "LINEABOVE" and cmd[1] == (0, 2) for cmd in style)
    assert data[2][-1].text == "Bebek koltuğu, çıkışta arayın"
    assert data[2][6].text == "05:20"
    assert data[3][-1] == "" and data[3][1].text == "AYŞE VELİ"

    outputs = export_driver_reports(path, SHEET, tmp_path / "out")
    pdfs = [p for p in outputs if p.suffix == ".pdf"]
    assert {p.name for p in pdfs} == {"Sofor_1.pdf", "Sofor_2.pdf"}
    for pdf in pdfs:
        assert pdf.read_bytes().startswith(b"%PDF")


def test_pdf_large_group_is_chunked_and_builds(tmp_path):
    path = _planning(tmp_path / "planning.xlsx", big=45)
    groups = driver_groups(read_blocks(path, SHEET))["3"]
    data, style = _pdf_table_parts(groups, "Helvetica", "Helvetica-Bold")
    assert len(data) == 1 + 45
    span_rows = sorted({cmd[1][1] for cmd in style if cmd[0] == "SPAN"})
    assert span_rows == [1, 21, 41]          # 20 + 20 + 5
    # Ortak alanlar her parçanın ilk satırında tekrarlanır.
    assert all(data[r][-1].text == "Büyük grup" for r in span_rows)
    # Aynı rezervasyonun parçaları arasında kalın rezervasyon çizgisi yok.
    assert not any(cmd[0] == "LINEABOVE" for cmd in style)
    outputs = export_driver_reports(path, SHEET, tmp_path / "out", selected_driver="3")
    pdf = next(p for p in outputs if p.suffix == ".pdf")
    raw = pdf.read_bytes()
    pages = raw.count(b"/Type /Page") - raw.count(b"/Type /Pages")
    assert pages >= 2


@pytest.mark.skipif(sys.platform != "win32", reason="Windows yazı tipleri")
def test_pdf_font_supports_turkish():
    regular, bold = _register_pdf_fonts()
    assert regular != "Helvetica" and bold != "Helvetica-Bold"


def test_hotel_card_name_region_and_meeting_point(tmp_path):
    from app.hotels import merge_names, resolver, upsert_card

    path = _planning(tmp_path / "planning.xlsx")
    catalog = tmp_path / "hotels.json"
    card = merge_names(["ARGOS"], name="ARGOS IN CAPPADOCIA", region="uçhisar", path=catalog)
    upsert_card({"id": card.id, "name": card.name, "meeting_point": "Kapı önü"}, path=catalog)
    groups = driver_groups(read_blocks(path, SHEET), resolver(catalog))["2"]
    argos = [g for g in groups if g["hotel"] == "ARGOS IN CAPPADOCIA"]
    assert argos and all(g["region"] == "Uçhisar" for g in argos)
    assert argos[0]["note"].startswith("Buluşma: Kapı önü")
    other = [g for g in groups if g["hotel"] != "ARGOS IN CAPPADOCIA"]
    assert all(g["region"] == "" for g in other)
