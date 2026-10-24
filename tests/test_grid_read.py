"""Hızlı okuma yolu (salt-okunur ızgara) tam yüklemeyle birebir aynı satırları vermeli."""
from __future__ import annotations

import shutil
import zipfile
from pathlib import Path

import openpyxl
import pytest

from app import config
from app.manifest import planning
from app.manifest.excel_cache import clear_excel_cache

FIXTURE = Path(__file__).parent / "fixtures" / "seat_template_ekim.xlsx"


@pytest.fixture()
def wb_path(tmp_path):
    p = tmp_path / "ekim.xlsx"
    shutil.copy2(FIXTURE, p)
    clear_excel_cache()
    return p


def test_grid_rows_match_full_load_on_every_sheet(wb_path):
    wb = openpyxl.load_workbook(wb_path, data_only=True)
    grid = planning._read_workbook_grid(wb_path)
    assert set(grid) == set(wb.sheetnames)
    for ws in wb.worksheets:
        assert planning._rows_from_grid(grid[ws.title]) == planning._rows_from_ws(ws), ws.title


def test_merged_ranges_are_read_from_xml(wb_path):
    wb = openpyxl.load_workbook(wb_path)
    expected = {ws.title: sorted(r.bounds for r in ws.merged_cells.ranges) for ws in wb.worksheets if ws.merged_cells.ranges}
    got = {k: sorted(v) for k, v in planning._merged_ranges(wb_path).items()}
    assert got == expected


def test_stale_value_under_merge_is_hidden(tmp_path):
    """Excel bazen birleşik alanın altındaki hücrede eski değer bırakır; okunmamalı."""
    p = tmp_path / "stale.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "01.11.2026"
    ws.cell(4, config.COL_PAX).value = 2
    ws.cell(4, config.COL_NAME).value = "A"
    ws.cell(4, config.COL_HOTEL).value = "OTEL-A"
    ws.cell(5, config.COL_NAME).value = "B"
    ws.cell(5, config.COL_HOTEL).value = "ESKI-DEGER"
    wb.save(p)
    # Birleşmeyi XML'e elle ekle: openpyxl merge_cells alttaki değeri sileceği için.
    col = openpyxl.utils.get_column_letter(config.COL_HOTEL)
    src = tmp_path / "src.xlsx"
    p.rename(src)
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(p, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"</sheetData>", f'</sheetData><mergeCells count="1"><mergeCell ref="{col}4:{col}5"/></mergeCells>'.encode())
            zout.writestr(item, data)
    clear_excel_cache()
    rows = planning.read_rows(p, "01.11.2026")
    assert [r.hotel for r in rows] == ["OTEL-A", ""]


def test_month_load_uses_shared_grid(wb_path, monkeypatch):
    calls = []
    real = planning._read_workbook_grid
    monkeypatch.setattr(planning, "_read_workbook_grid", lambda p: calls.append(p) or real(p))
    clear_excel_cache()
    planning.read_rows(wb_path, "22.10.2026")
    planning.read_rows(wb_path, "23.10.2026")
    planning.month_load(wb_path, 2026, 10, capacity=28, codes=["BYF", "BTK", "BYJ", "BZR", "BZV"])
    assert len(calls) == 1   # üç okuma, tek kitap açılışı
