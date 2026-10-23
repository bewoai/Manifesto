"""Uygulama içi Excel görünümü: gün sayfasını okuma ve hücre düzenleme."""
from __future__ import annotations

import importlib
import json
import shutil
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import config

FIXTURE = Path(__file__).parent / "fixtures" / "seat_template_ekim.xlsx"


@pytest.fixture()
def env(monkeypatch, tmp_path):
    plan = tmp_path / "plan.xlsx"
    shutil.copy2(FIXTURE, plan)
    settings_path = tmp_path / "settings.json"
    settings_path.write_text(json.dumps({"planning_xlsx": str(plan), "output_dir": str(tmp_path / "out")}), encoding="utf-8")
    monkeypatch.setenv("MANIFESTO_SETTINGS", str(settings_path))
    monkeypatch.setenv("IRTIFA_DB_PATH", str(tmp_path / "v.db"))
    import app.settings as settings_mod
    importlib.reload(settings_mod)
    import app.main as main_mod
    importlib.reload(main_mod)
    from app.manifest.excel_cache import clear_excel_cache
    clear_excel_cache()
    with TestClient(main_mod.app) as client:
        yield client, plan


def test_view_returns_real_row_numbers_and_header(env):
    client, plan = env
    data = client.get("/api/planning/view", params={"sheet": "22.10.2026"}).json()
    assert data["header_row"] == config.PLANNING_HEADER_ROW and data["first_data_row"] == config.PLANNING_FIRST_DATA_ROW
    assert data["header"] and data["workbook_revision"]
    ws = openpyxl.load_workbook(plan, data_only=True)["22.10.2026"]
    first = data["rows"][0]
    assert first["row"] >= config.PLANNING_FIRST_DATA_ROW
    name = ws.cell(first["row"], config.COL_NAME).value
    assert first["cells"][config.COL_NAME - 1] == ("" if name is None else str(name).strip())


def test_view_unknown_sheet_404(env):
    client, _ = env
    assert client.get("/api/planning/view", params={"sheet": "01.01.2030"}).status_code == 404


def test_cells_write_batch_and_keeps_text_passport(env):
    client, plan = env
    view = client.get("/api/planning/view", params={"sheet": "22.10.2026"}).json()
    target = next(r for r in view["rows"] if r["cells"][config.COL_NAME - 1])["row"]
    r = client.post("/api/planning/cells", json={
        "sheet": "22.10.2026", "expected_revision": view["workbook_revision"],
        "changes": [
            {"row": target, "col": config.COL_PICKUP, "value": "04:30"},
            {"row": target, "col": config.COL_PASSPORT_NO, "value": "01234567"},
            {"row": target, "col": config.COL_NAME, "value": "yeni isim"},
        ],
    })
    assert r.status_code == 200, r.text
    assert r.json()["count"] == 3 and r.json()["backup"]
    ws = openpyxl.load_workbook(plan)["22.10.2026"]
    assert ws.cell(target, config.COL_PICKUP).value == "04:30"
    cell = ws.cell(target, config.COL_PASSPORT_NO)
    assert cell.value == "01234567" and cell.number_format == "@"
    assert ws.cell(target, config.COL_NAME).value == "YENI ISIM"
    # Görünüm de yeni değeri okur (önbellek dosya değişince tazelenir)
    again = client.get("/api/planning/view", params={"sheet": "22.10.2026"}).json()
    assert next(x for x in again["rows"] if x["row"] == target)["cells"][config.COL_PICKUP - 1] == "04:30"


def test_cells_reject_stale_revision(env):
    client, _ = env
    view = client.get("/api/planning/view", params={"sheet": "22.10.2026"}).json()
    row = view["rows"][0]["row"]
    ok = client.post("/api/planning/cells", json={"sheet": "22.10.2026", "expected_revision": view["workbook_revision"],
                                                  "changes": [{"row": row, "col": config.COL_NOTE, "value": "a"}]})
    assert ok.status_code == 200
    stale = client.post("/api/planning/cells", json={"sheet": "22.10.2026", "expected_revision": view["workbook_revision"],
                                                     "changes": [{"row": row, "col": config.COL_NOTE, "value": "b"}]})
    assert stale.status_code == 409


@pytest.mark.parametrize("change", [
    {"row": 2, "col": 1, "value": "x"},                       # başlık satırı
    {"row": 10, "col": config.COL_PASSPORT_NO + 1, "value": "x"},
    {"row": 10, "col": config.COL_PAX, "value": "iki"},       # PAX sayı değil
    {"row": "a", "col": 1, "value": "x"},
])
def test_cells_validation(env, change):
    client, _ = env
    r = client.post("/api/planning/cells", json={"sheet": "22.10.2026", "changes": [change]})
    assert r.status_code == 422


def test_cells_refuse_inside_merged_range(env):
    client, plan = env
    wb = openpyxl.load_workbook(plan)
    ws = wb["22.10.2026"]
    rng = next((m for m in ws.merged_cells.ranges if m.min_row >= config.PLANNING_FIRST_DATA_ROW and m.max_row > m.min_row), None)
    if rng is None:
        pytest.skip("fixture'da veri satırında birleşik alan yok")
    r = client.post("/api/planning/cells", json={"sheet": "22.10.2026",
                                                 "changes": [{"row": rng.min_row + 1, "col": rng.min_col, "value": "x"}]})
    assert r.status_code == 422 and "birleşik" in r.json()["detail"]
