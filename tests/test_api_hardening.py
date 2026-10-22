"""API sertleştirme: manifesto export'ta balon kodu doğrulaması (path traversal) + CORS."""
from __future__ import annotations

import importlib

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import config


@pytest.fixture()
def client(monkeypatch, tmp_path):
    monkeypatch.setenv("MANIFESTO_SETTINGS", str(tmp_path / "settings.json"))
    monkeypatch.setenv("IRTIFA_DB_PATH", str(tmp_path / "hardening.db"))
    import app.settings as settings_mod
    importlib.reload(settings_mod)
    import app.main as main_mod
    importlib.reload(main_mod)
    return TestClient(main_mod.app)


# ─── Balon kodu doğrulaması ──────────────────────────────────────────────

def test_normalize_balloon_code_accepts_real_codes():
    from app.manifest.writer import normalize_balloon_code

    assert normalize_balloon_code(" bzr ") == "BZR"
    assert normalize_balloon_code("byf-2") == "BYF-2"
    assert normalize_balloon_code("B_01") == "B_01"


@pytest.mark.parametrize("bad", ["../x", "..\\..\\x", "BZR/..", "BZR.xlsx", "", "   ", "C:\\tmp\\x", "-BZR", "A" * 17])
def test_normalize_balloon_code_rejects_unsafe(bad):
    from app.manifest.writer import normalize_balloon_code

    with pytest.raises(ValueError):
        normalize_balloon_code(bad)


@pytest.mark.parametrize("bad", ["../x", "..%5C..%5Cx", "BZR/..", "BZR.xlsx", ""])
@pytest.mark.parametrize("endpoint", ["/api/manifest/export", "/api/manifest/preview", "/manifest/export", "/manifest/preview"])
def test_manifest_endpoints_reject_unsafe_balloon(client, endpoint, bad):
    key = "date" if endpoint.startswith("/manifest/") else "sheet"
    r = client.get(endpoint, params={key: "22.06.2026", "balloon": bad})
    assert r.status_code == 422, (endpoint, bad, r.status_code, r.text)


def _planning_workbook(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "22.06.2026"
    ws.cell(4, config.COL_PAX).value = 1
    ws.cell(4, config.COL_NAME).value = "TEST PERSON"
    ws.cell(4, config.COL_UYRUK).value = "TUR"
    ws.cell(4, config.COL_MF).value = "F"
    ws.cell(4, config.COL_PASSPORT_NO).value = "U0001"
    ws.cell(4, config.COL_BALLOON).value = "BZR"
    wb.save(path)
    wb.close()


def test_export_writes_only_inside_out_dir(tmp_path, template_ready):
    from app.manifest.writer import export

    planning = tmp_path / "planning.xlsx"
    _planning_workbook(planning)
    out_dir = tmp_path / "out"

    path = export(planning, "22.06.2026", "bzr", out_dir)
    assert path == out_dir / "BZR.xlsx" and path.exists()

    # Sayfadan gelen tuhaf bir kod bile dosya adında temizlenir, klasör dışına çıkamaz
    odd = export(planning, "22.06.2026", "../BZR", out_dir)
    assert odd.parent == out_dir
    assert odd.name == "___BZR.xlsx"
    assert not (tmp_path / "BZR.xlsx").exists()


# ─── CORS ────────────────────────────────────────────────────────────────

def test_cors_blocks_foreign_origin(client):
    r = client.get("/health", headers={"Origin": "https://evil.example"})
    assert r.status_code == 200
    assert "access-control-allow-origin" not in r.headers

    pre = client.options(
        "/api/settings",
        headers={"Origin": "https://evil.example", "Access-Control-Request-Method": "GET"},
    )
    assert "access-control-allow-origin" not in pre.headers


@pytest.mark.parametrize("origin", ["http://127.0.0.1:5173", "http://localhost:8000", "http://127.0.0.1"])
def test_cors_allows_local_origins(client, origin):
    r = client.get("/health", headers={"Origin": origin})
    assert r.headers.get("access-control-allow-origin") == origin
