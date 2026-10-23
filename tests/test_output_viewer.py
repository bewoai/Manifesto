"""Çıktı dosyalarını uygulama içinde listeleme ve Excel görüntüleme."""
from __future__ import annotations

import importlib
import json

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import config


@pytest.fixture()
def env(monkeypatch, tmp_path):
    out = tmp_path / "out"
    plan = tmp_path / "plan.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "22.10.2026"
    ws.cell(4, config.COL_PAX).value = 1
    ws.cell(4, config.COL_NAME).value = "TEST KISI"
    ws.cell(4, config.COL_UYRUK).value = "TUR"
    ws.cell(4, config.COL_MF).value = "F"
    ws.cell(4, config.COL_PASSPORT_NO).value = "01234567"
    ws.cell(4, config.COL_BALLOON).value = "BZR"
    wb.save(plan)
    settings_path = tmp_path / "settings.json"
    settings_path.write_text(json.dumps({"planning_xlsx": str(plan), "output_dir": str(out)}), encoding="utf-8")
    monkeypatch.setenv("MANIFESTO_SETTINGS", str(settings_path))
    monkeypatch.setenv("IRTIFA_DB_PATH", str(tmp_path / "v.db"))
    import app.settings as settings_mod
    importlib.reload(settings_mod)
    import app.main as main_mod
    importlib.reload(main_mod)
    from app.manifest.excel_cache import clear_excel_cache
    clear_excel_cache()
    with TestClient(main_mod.app) as client:
        yield client, out


def test_single_export_goes_to_day_folder_and_is_listed(env):
    client, out = env
    r = client.get("/api/manifest/export", params={"sheet": "22.10.2026", "balloon": "BZR", "override_reason": "test"})
    assert r.status_code == 200, r.text
    assert (out / "22.10.2026" / "Manifestolar" / "BZR.xlsx").is_file()
    files = client.get("/api/outputs/files", params={"sheet": "22.10.2026"}).json()["files"]
    assert [(f["path"], f["group"], f["kind"]) for f in files] == [("Manifestolar/BZR.xlsx", "Manifestolar", "excel")]


def test_xlsx_view_returns_manifest_sheet_first_with_text_passport(env):
    client, _ = env
    client.get("/api/manifest/export", params={"sheet": "22.10.2026", "balloon": "BZR", "override_reason": "test"})
    data = client.get("/api/outputs/xlsx", params={"sheet": "22.10.2026", "path": "Manifestolar/BZR.xlsx"}).json()
    first = data["sheets"][0]
    assert first["name"] == config.MANIFEST_SHEET
    row = first["rows"][config.MANIFEST_FIRST_DATA_ROW - 1]
    assert row[config.M_COL_NAME - 1] == "TEST KISI"
    assert row[config.M_COL_PASSPORT - 1] == "01234567"


@pytest.mark.parametrize("path", ["../plan.xlsx", "../../x.xlsx", "Manifestolar/../../plan.xlsx", "Manifestolar/yok.xlsx", "a.exe"])
def test_viewer_rejects_paths_outside_day_folder(env, path):
    client, _ = env
    client.get("/api/manifest/export", params={"sheet": "22.10.2026", "balloon": "BZR", "override_reason": "test"})
    r = client.get("/api/outputs/xlsx", params={"sheet": "22.10.2026", "path": path})
    assert r.status_code in (404, 422)
    r = client.post("/api/outputs/open-file", json={"sheet": "22.10.2026", "path": path})
    assert r.status_code in (404, 422)


def test_files_for_day_without_outputs_is_empty(env):
    client, _ = env
    assert client.get("/api/outputs/files", params={"sheet": "23.10.2026"}).json()["files"] == []
