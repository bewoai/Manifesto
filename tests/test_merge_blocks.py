"""Rezervasyonların ortak kolonlarını birleşik düzene getirme (PAX 4 doğru örnek gibi)."""
from __future__ import annotations

import datetime as dt
import importlib
import json

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import config
from app.manifest.excel_cache import clear_excel_cache
from app.manifest.planning import merge_reservation_blocks, preview_block_merges, read_blocks


def _repeated_block_workbook(path):
    """PAX 6 yanlış örnek: ortak bilgiler her satırda tekrar yazılmış; ikinci blokta otel farklı."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "24.10.2026"
    ws.cell(config.PLANNING_HEADER_ROW, config.COL_PAX).value = "PAX"
    for i in range(3):
        r = 4 + i
        ws.cell(r, config.COL_PAX).value = 3 if i == 0 else None
        ws.cell(r, config.COL_NAME).value = f"KISI {i}"
        ws.cell(r, config.COL_ROOM).value = "0 532 411 97 59"
        ws.cell(r, config.COL_HOTEL).value = "MARRIOT"
        ws.cell(r, config.COL_PICKUP).value = dt.time(5, 0)
        ws.cell(r, config.COL_RESERVED_BY).value = "MAHMUT"
        ws.cell(r, config.COL_AGENCY).value = "BÜLENT KUNAÇ"
        ws.cell(r, config.COL_BALLOON).value = "BYF"
    for i in range(2):
        r = 7 + i
        ws.cell(r, config.COL_PAX).value = 2 if i == 0 else None
        ws.cell(r, config.COL_NAME).value = f"YOLCU {i}"
        ws.cell(r, config.COL_HOTEL).value = "ARGOS" if i == 0 else "KELEBEK"   # çelişki → dokunulmaz
        ws.cell(r, config.COL_AGENCY).value = "KARAVAN" if i == 0 else None
        ws.cell(r, config.COL_BALLOON).value = "BYF"
    wb.save(path)


def test_merge_repeated_block_keeps_values_and_skips_conflicts(tmp_path):
    p = tmp_path / "plan.xlsx"
    _repeated_block_workbook(p)
    clear_excel_cache()
    preview = preview_block_merges(p, "24.10.2026")
    assert preview["blocks"] == 2
    result = merge_reservation_blocks(p, "24.10.2026")
    assert result["blocks"] == 2
    assert any(sk["lead_row"] == 7 and sk["col"] == config.COL_HOTEL for sk in result["skipped"])

    ws = openpyxl.load_workbook(p)["24.10.2026"]
    ranges = {(m.min_col, m.min_row, m.max_col, m.max_row) for m in ws.merged_cells.ranges}
    for col in (config.COL_PAX, config.COL_HOTEL, config.COL_PICKUP, config.COL_RESERVED_BY, config.COL_AGENCY):
        assert (col, 4, col, 6) in ranges
    # Sayfada örnek yokken irtibat/oda ve not birleştirilmez (yolcu başına yazılabilir)
    assert (config.COL_ROOM, 4, config.COL_ROOM, 6) not in ranges
    assert ws.cell(5, config.COL_ROOM).value == "0 532 411 97 59"
    assert ws.cell(4, config.COL_PICKUP).value == dt.time(5, 0)     # tip korunur
    assert ws.cell(4, config.COL_HOTEL).value == "MARRIOT" and ws.cell(4, config.COL_PAX).value == 3
    assert (config.COL_AGENCY, 7, config.COL_AGENCY, 8) in ranges    # tek dolu değer → birleşir
    assert (config.COL_HOTEL, 7, config.COL_HOTEL, 8) not in ranges  # iki farklı otel → dokunulmaz
    assert ws.cell(8, config.COL_HOTEL).value == "KELEBEK"
    assert ws.cell(4, config.COL_BALLOON).value == "BYF" and ws.cell(5, config.COL_BALLOON).value == "BYF"  # balon birleştirilmez

    clear_excel_cache()
    blocks = read_blocks(p, "24.10.2026")
    assert [len(b.rows) for b in blocks] == [3, 2]
    assert blocks[0].hotel == "MARRIOT" and blocks[0].agency == "BÜLENT KUNAÇ"
    # İkinci kez çalıştırmak bir şey değiştirmez
    assert preview_block_merges(p, "24.10.2026")["blocks"] == 0


@pytest.fixture()
def client(monkeypatch, tmp_path):
    plan = tmp_path / "plan.xlsx"
    _repeated_block_workbook(plan)
    settings_path = tmp_path / "settings.json"
    settings_path.write_text(json.dumps({"planning_xlsx": str(plan)}), encoding="utf-8")
    monkeypatch.setenv("MANIFESTO_SETTINGS", str(settings_path))
    monkeypatch.setenv("IRTIFA_DB_PATH", str(tmp_path / "m.db"))
    import app.settings as settings_mod
    importlib.reload(settings_mod)
    import app.main as main_mod
    importlib.reload(main_mod)
    clear_excel_cache()
    with TestClient(main_mod.app) as c:
        yield c


def test_view_reports_merges_and_plan_then_endpoint_applies(client):
    view = client.get("/api/planning/view", params={"sheet": "24.10.2026"}).json()
    assert view["merges"] == [] and view["merge_plan"]["blocks"] == 2
    r = client.post("/api/planning/merge-blocks", json={"sheet": "24.10.2026", "expected_revision": view["workbook_revision"]})
    assert r.status_code == 200, r.text
    assert r.json()["blocks"] == 2 and r.json()["backup"]
    view2 = client.get("/api/planning/view", params={"sheet": "24.10.2026"}).json()
    assert {"row": 4, "col": config.COL_HOTEL, "rowspan": 3, "colspan": 1} in view2["merges"]
    assert view2["merge_plan"]["blocks"] == 0



def test_merge_columns_follow_the_sheets_own_convention(tmp_path):
    """24.10 gerçek düzeni: A/F/G/H/I her rezervasyonda birleşik, E ve M yalnız birkaçında.
    Düzeltme yalnız sayfanın çoğunlukla birleştirdiği kolonları birleştirir."""
    from app.manifest.planning import block_merge_columns, read_rows

    p = tmp_path / "conv.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "24.10.2026"
    core = (config.COL_PAX, config.COL_HOTEL, config.COL_PICKUP, config.COL_RESERVED_BY, config.COL_AGENCY)
    lead = 4
    for i in range(4):                       # 4 rezervasyon × 2 yolcu
        for k in range(2):
            r = lead + k
            ws.cell(r, config.COL_NAME).value = f"K{i}{k}"
            ws.cell(r, config.COL_BALLOON).value = "BYF"
            ws.cell(r, config.COL_ROOM).value = f"ODA {i}{k}"     # yolcu başına oda
            ws.cell(r, config.COL_NOTE).value = "VEGAN" if k == 0 else None
        ws.cell(lead, config.COL_PAX).value = 2
        ws.cell(lead, config.COL_HOTEL).value = "ARGOS"
        if i < 3:                             # ilk 3 rezervasyon düzende, sonuncusu değil
            for c in core:
                ws.merge_cells(start_row=lead, start_column=c, end_row=lead + 1, end_column=c)
        if i == 0:                            # E yalnız 1 rezervasyonda birleşik
            ws.cell(lead + 1, config.COL_ROOM).value = None
            ws.merge_cells(start_row=lead, start_column=config.COL_ROOM, end_row=lead + 1, end_column=config.COL_ROOM)
        lead += 2
    wb.save(p)
    clear_excel_cache()
    merged = [(m.min_col, m.min_row, m.max_col, m.max_row) for m in openpyxl.load_workbook(p)["24.10.2026"].merged_cells.ranges]
    assert block_merge_columns(read_rows(p, "24.10.2026"), merged) == core

    preview = preview_block_merges(p, "24.10.2026")
    assert preview["blocks"] == 1 and {a["lead_row"] for a in preview["actions"]} == {10}
    assert {a["col"] for a in preview["actions"]} <= set(core)
    merge_reservation_blocks(p, "24.10.2026")
    ws2 = openpyxl.load_workbook(p)["24.10.2026"]
    assert ws2.cell(11, config.COL_ROOM).value == "ODA 31"      # oda yolcu başına kaldı
    assert ws2.cell(10, config.COL_NOTE).value == "VEGAN"         # not birleştirilmedi
    assert not any(m.min_col == config.COL_NOTE for m in ws2.merged_cells.ranges)
